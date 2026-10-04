import shutil
import unittest
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from agent_hr import ROOT
from application_preparer import build_cover_letter
from cv_generator import generate_tailored_cv
from document_outputs import docx_to_pdf, text_to_pdf
from document_quality import build_quality_report
from job_analysis import analyse_job
from source_truth import verify_source_truth


CASES = {
    "operational_risk": {
        "id": 9911,
        "title": "Operational Risk Management (ORM) Specialist (m/f/d) - career A1",
        "company": "Banque centrale du Luxembourg",
        "location": "Luxembourg",
        "description": (
            "Operational risk assessments, risk matrices, internal controls, DORA, governance and remediation plans. "
            "Fluent English required; French or German is an asset."
        ),
    },
    "aml_junior": {
        "id": 9912,
        "title": "Junior AML/KYC Analyst",
        "company": "Example Private Bank",
        "location": "Luxembourg",
        "description": (
            "Junior role supporting AML/KYC reviews, customer due diligence and transaction monitoring. "
            "English and French required. Training provided."
        ),
    },
    "change": {
        "id": 9913,
        "title": "Change Management Consultant",
        "company": "Transformation Partners",
        "location": "Paris",
        "description": (
            "Change management role covering stakeholder workshops, user adoption, training, operating model, "
            "PMO, project monitoring and digital transformation. English and French required."
        ),
    },
}


class CandidateGenerationTests(unittest.TestCase):
    def tearDown(self):
        for case in CASES.values():
            shutil.rmtree(ROOT / "applications" / f"application_{case['id']:04d}", ignore_errors=True)

    def _generate(self, case):
        analysis = analyse_job(case["title"], case["company"], case["location"], case["description"])
        docx_path = generate_tailored_cv(
            case["id"], case["title"], case["company"], case["location"], case["description"], analysis=analysis
        )
        pdf_path = docx_to_pdf(docx_path)
        letter_text = build_cover_letter(
            case["title"], case["company"], case["location"], case["description"]
        )
        letter_path = text_to_pdf(
            letter_text,
            docx_path.parent / f"Demo_Candidate_Cover_Letter_{case['company'].replace(' ', '_')}.pdf",
        )
        report = build_quality_report(pdf_path, letter_path, analysis=analysis, docx_path=docx_path)
        return analysis, docx_path, pdf_path, letter_path, report

    def test_immutable_source_cv_checksum_is_verified(self):
        profile = verify_source_truth()
        self.assertEqual(
            profile["source_truth"]["sha256"],
            "7369740FB2D7E04D1D0CAD5630141ED5FA962225B73E55A53AD77AF1862EADE0",
        )

    def test_multidimensional_matching_does_not_hide_seniority_gap(self):
        analysis = analyse_job(
            "AML Analyst - Senior",
            "EY Luxembourg",
            "Luxembourg",
            "Minimum 5 years of relevant AML experience required. Transaction monitoring and enhanced due diligence. Fluent English required.",
        )
        self.assertLess(analysis["score"], 45)
        self.assertFalse(analysis["eligible_for_proposal"])
        self.assertLessEqual(analysis["dimensions"]["seniority"], 5)
        self.assertEqual(analysis["dimensions"]["domain_experience"], 0)

    def test_three_representative_offers_create_distinct_source_grounded_documents(self):
        outputs = {name: self._generate(case) for name, case in CASES.items()}
        fingerprints = set()
        for name, (analysis, docx_path, pdf_path, letter_path, report) in outputs.items():
            self.assertTrue(report["passed"], f"{name}: {report['errors']}")
            self.assertIn(len(PdfReader(str(pdf_path)).pages), (1, 2))
            self.assertEqual(len(PdfReader(str(letter_path)).pages), 1)
            self.assertNotRegex(pdf_path.name.lower(), r"ats|optimized|generated|tailored|match")
            self.assertNotRegex(letter_path.name.lower(), r"ats|optimized|generated|tailored|match")
            fingerprints.add(report["fingerprint"])
            document = Document(docx_path)
            self.assertEqual(len(document.tables), 0)
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)
            self.assertIn("Qualitel", text)
            self.assertIn("Embassy of Qatar to UNESCO", text)
            self.assertNotIn("m/f/d", text)
            self.assertNotIn("career A1", text)
            self.assertNotIn("ATS", text)
        self.assertEqual(len(fingerprints), 3)

        operational_text = "\n".join(p.text for p in Document(outputs["operational_risk"][1]).paragraphs)
        change_text = "\n".join(p.text for p in Document(outputs["change"][1]).paragraphs)
        aml_text = "\n".join(p.text for p in Document(outputs["aml_junior"][1]).paragraphs)
        self.assertIn("Risk, Resilience and Compliance Consultant", operational_text)
        self.assertIn("Change Management and Digital Transformation Consultant", change_text)
        self.assertIn("AML/KYC training in progress", aml_text)
        self.assertNotIn("AML expert", aml_text)
        self.assertNotEqual(operational_text, change_text)


if __name__ == "__main__":
    unittest.main()
