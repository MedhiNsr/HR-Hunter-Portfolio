import re

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from agent_hr import ROOT
from job_analysis import analyse_job, infer_role_family
from job_intelligence import clean_job_title
from source_truth import verify_source_truth


NAVY = RGBColor(24, 50, 74)
STEEL = RGBColor(73, 91, 108)
TEXT = RGBColor(32, 36, 40)

FAMILY_HEADLINES = {
    "operational_risk": "Risk, Resilience and Compliance Consultant",
    "aml_kyc": "Risk and Compliance Consultant",
    "change_transformation": "Change Management and Digital Transformation Consultant",
    "compliance_risk": "Regulatory Risk and Compliance Consultant",
}

FAMILY_TERMS = {
    "operational_risk": ["dora", "operational", "resilience", "risk", "control", "governance", "remediation", "nis2", "iso 27001", "banking", "insurance"],
    "aml_kyc": ["regulatory", "compliance", "control", "risk", "banking", "insurance", "monitoring", "governance", "documentation"],
    "change_transformation": ["change", "transformation", "deployment", "training", "user", "procedure", "pmo", "roadmap", "operating model", "stakeholder", "implementation"],
    "compliance_risk": ["regulatory", "compliance", "control", "risk", "governance", "monitoring", "assessment", "stakeholder", "documentation"],
}

SKILL_LABELS = {
    "DORA": "Digital Operational Resilience Act (DORA)",
    "Operational resilience": "Operational resilience",
    "Operational risk": "Operational risk assessment",
    "Internal controls": "Internal controls and control frameworks",
    "Regulatory monitoring": "Regulatory monitoring",
    "Governance": "Risk and project governance",
    "Change management": "Change management and user adoption",
    "Digital transformation": "Digital transformation",
    "Project coordination": "Project and PMO coordination",
    "Training and workshops": "Workshops, training and procedures",
    "Stakeholder management": "Stakeholder coordination",
}


def load_content():
    return verify_source_truth()


def header_location(location):
    text = (location or "").lower()
    if "luxembourg" in text:
        return "Luxembourg"
    if "paris" in text or "france" in text:
        return "Paris"
    return "Paris / Luxembourg"


def role_family(title, description):
    return infer_role_family(title, description)


def _natural_join(items):
    items = [item for item in items if item]
    if len(items) < 2:
        return items[0] if items else ""
    return ", ".join(items[:-1]) + " and " + items[-1]


def tailored_summary(title, company, location, description, analysis=None):
    analysis = analysis or analyse_job(title, company, location, description)
    family = analysis["role_family"]
    supported = [SKILL_LABELS[item] for item in analysis["supported_keywords"] if item in SKILL_LABELS][:3]
    if not supported:
        supported = {
            "operational_risk": ["DORA", "operational resilience", "internal controls"],
            "aml_kyc": ["regulatory monitoring", "internal controls", "risk assessment"],
            "change_transformation": ["change management", "digital transformation", "stakeholder coordination"],
            "compliance_risk": ["regulatory monitoring", "internal controls", "risk assessment"],
        }[family]
    opening = {
        "operational_risk": "Risk, resilience and compliance consultant with experience supporting banks and insurers",
        "aml_kyc": "Early-career risk and compliance consultant with financial-services exposure",
        "change_transformation": "Change and digital-transformation consultant experienced in complex delivery environments",
        "compliance_risk": "Regulatory risk and compliance consultant with experience in regulated environments",
    }[family]
    summary = f"{opening}, with hands-on work in {_natural_join(supported)}."
    if family == "aml_kyc":
        summary += " Currently completing specialist AML/KYC training, supported by transferable regulatory and control work."
    else:
        summary += " Brings structured analysis, clear documentation and practical stakeholder coordination."
    return summary + f" Native French speaker, fluent in English and Spanish, and available immediately in {header_location(location)}."


def tailored_skills(title, description, analysis=None):
    analysis = analysis or analyse_job(title, "", "", description)
    skills = []
    for keyword in analysis["supported_keywords"]:
        label = SKILL_LABELS.get(keyword)
        if label and label not in skills:
            skills.append(label)
    fallbacks = {
        "operational_risk": ["Risk matrices and remediation roadmaps", "NIS2 and ISO 27001 regulatory monitoring", "Process documentation", "Stakeholder workshops and training"],
        "aml_kyc": ["Regulatory requirements analysis", "Internal controls", "Risk assessment", "Financial-services consulting", "Process documentation", "AML/KYC training in progress"],
        "change_transformation": ["Change management and user adoption", "Digital transformation", "Project and PMO coordination", "Operating-model support", "Training and operational procedures", "Stakeholder coordination"],
        "compliance_risk": ["Regulatory compliance", "Internal controls", "Risk assessment and remediation planning", "Governance documentation", "Stakeholder coordination", "Executive-ready synthesis"],
    }[analysis["role_family"]]
    for label in fallbacks:
        if label not in skills:
            skills.append(label)
    return skills[:8]


def _add_style(document, name, size, bold=False, color=TEXT, before=0, after=0, leading=1.08):
    styles = document.styles
    style = styles[name] if name in styles else styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Arial"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = leading
    return style


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Cm(1.25)
    section.bottom_margin = Cm(1.2)
    section.left_margin = Cm(1.45)
    section.right_margin = Cm(1.45)
    _add_style(document, "CV Name", 20, True, NAVY, after=2)
    _add_style(document, "CV Headline", 10.5, True, STEEL, after=3)
    _add_style(document, "CV Contact", 8.8, False, STEEL, after=1)
    _add_style(document, "CV Section", 10.5, True, NAVY, before=7, after=3)
    _add_style(document, "CV Job", 9.6, True, NAVY, before=4, after=1)
    _add_style(document, "CV Job Meta", 8.7, True, STEEL, after=2)
    _add_style(document, "CV Page Break", 1, False, TEXT)
    _add_style(document, "CV Body", 9.1, False, TEXT, after=2, leading=1.12)
    bullet = _add_style(document, "CV Bullet", 8.9, False, TEXT, after=1.8, leading=1.1)
    bullet.paragraph_format.left_indent = Cm(0.38)
    bullet.paragraph_format.first_line_indent = Cm(-0.22)


def add_bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "5")
    bottom.set(qn("w:color"), "B8C2CC")
    borders.append(bottom)
    p_pr.append(borders)


def add_bullets(document, bullets):
    for bullet in bullets:
        document.add_paragraph(f"- {bullet}", style="CV Bullet")


def _bullet_score(bullet, analysis):
    lower = bullet.lower()
    terms = list(FAMILY_TERMS[analysis["role_family"]])
    terms.extend(keyword.lower() for keyword in analysis["supported_keywords"])
    score = sum(3 for term in terms if term in lower)
    if analysis["sector"] == "financial_services" and any(term in lower for term in ["bank", "insurance", "financial"]):
        score += 2
    return score


def experience_bullets(experience, analysis):
    limits = {
        "operational_risk": {"EURUS": 7, "AP-HP": 3, "DXC Technology for Agirc-Arrco": 3, "Qualitel": 2, "Embassy of Qatar to UNESCO": 2},
        "aml_kyc": {"EURUS": 6, "AP-HP": 3, "DXC Technology for Agirc-Arrco": 3, "Qualitel": 2, "Embassy of Qatar to UNESCO": 2},
        "change_transformation": {"EURUS": 4, "AP-HP": 5, "DXC Technology for Agirc-Arrco": 5, "Qualitel": 3, "Embassy of Qatar to UNESCO": 2},
        "compliance_risk": {"EURUS": 7, "AP-HP": 3, "DXC Technology for Agirc-Arrco": 3, "Qualitel": 2, "Embassy of Qatar to UNESCO": 2},
    }[analysis["role_family"]]
    ranked = sorted(enumerate(experience["bullets"]), key=lambda item: (-_bullet_score(item[1], analysis), item[0]))
    return [bullet for _, bullet in ranked[:limits.get(experience["company"], 3)]]


def _safe_company_name(company):
    return (re.sub(r"[^A-Za-z0-9]+", "_", company or "Company").strip("_")[:45] or "Company")


def generate_tailored_cv(application_id, title, company, location, description, analysis=None):
    analysis = analysis or analyse_job(title, company, location, description)
    folder = ROOT / "applications" / f"application_{application_id:04d}"
    folder.mkdir(parents=True, exist_ok=True)
    path = folder / f"Demo_Candidate_CV_{_safe_company_name(company)}.docx"
    content = load_content()
    document = Document()
    configure_document(document)

    name = document.add_paragraph("Demo Candidate", style="CV Name")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    headline = document.add_paragraph(FAMILY_HEADLINES[analysis["role_family"]], style="CV Headline")
    headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact = document.add_paragraph(f"{header_location(location)} | +352 000 00 00 00 | candidate@example.com", style="CV Contact")
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    linkedin = document.add_paragraph("linkedin.com/in/demo-candidate", style="CV Contact")
    linkedin.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_bottom_border(linkedin)

    document.add_paragraph("Profile", style="CV Section")
    document.add_paragraph(tailored_summary(title, company, location, description, analysis), style="CV Body")
    document.add_paragraph("Core Expertise", style="CV Section")
    skills = tailored_skills(title, description, analysis)
    document.add_paragraph(" | ".join(skills[:4]), style="CV Body")
    document.add_paragraph(" | ".join(skills[4:]), style="CV Body")

    document.add_paragraph("Professional Experience", style="CV Section")
    for experience in content["experiences"]:
        if experience["company"] == "DXC Technology for Agirc-Arrco":
            page_break = document.add_paragraph(style="CV Page Break")
            page_break.add_run().add_break(WD_BREAK.PAGE)
        document.add_paragraph(experience["role"], style="CV Job")
        document.add_paragraph(f"{experience['company']} | {experience['location']} | {experience['dates']}", style="CV Job Meta")
        add_bullets(document, experience_bullets(experience, analysis))

    document.add_paragraph("Education", style="CV Section")
    add_bullets(document, content["education"])
    document.add_paragraph("Certifications", style="CV Section")
    add_bullets(document, content["certifications"])
    document.add_paragraph("Languages", style="CV Section")
    document.add_paragraph("French: native | English: fluent (C2) | Spanish: fluent (C2) | Literary Arabic: basic (A2)", style="CV Body")
    document.add_paragraph("Affiliations", style="CV Section")
    document.add_paragraph(" | ".join(content["affiliations"]), style="CV Body")

    document.core_properties.title = "Demo Candidate - Curriculum Vitae"
    document.core_properties.subject = clean_job_title(title)
    document.core_properties.author = "Demo Candidate"
    document.save(path)
    return path
