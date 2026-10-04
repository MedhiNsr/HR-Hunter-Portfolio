import hashlib
import json
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent
PROFILE_PATH = ROOT / "profile" / "candidate_cv_content.json"

INTERNAL_TERMS = (
    "ats", "optimized", "optimised", "generated", "ai generated", "tailored", "prompt",
    "hr hunter", "match score", "application message", "role-relevant themes", "career a1",
    "m/f/d", "m/f/x", "m/f/t", "h/f",
)
UNSUPPORTED_CLAIMS = (
    "suspicious activity investigations", "transaction monitoring experience", "sanctions screening experience",
    "enhanced due diligence experience", "sar filing", "str filing", "python", "sql", "people management",
    "senior aml expert", "five years of aml experience",
)
GENERIC_AI_PHRASES = (
    "i am thrilled to apply", "i am excited to apply", "perfect fit", "unique opportunity",
    "admired your company for", "passionate about joining", "dynamic team",
)
REQUIRED_CV_MARKERS = (
    "eurus", "ap-hp", "dxc technology", "qualitel", "embassy of qatar",
    "education", "certifications", "languages",
)


def extract_pdf_pages(path):
    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    return pages


def _source_bullets():
    profile = json.loads(PROFILE_PATH.read_text(encoding="utf-8"))
    allowed = set(profile["education"] + profile["certifications"])
    for experience in profile["experiences"]:
        allowed.update(experience["bullets"])
    return allowed


def _contains_term(text, term):
    return bool(re.search(rf"(?<!\w){re.escape(term)}(?!\w)", text, flags=re.IGNORECASE))


def build_quality_report(cv_path, cover_letter_path, analysis=None, docx_path=None):
    cv_path = Path(cv_path)
    cover_letter_path = Path(cover_letter_path)
    errors = []
    warnings = []
    try:
        cv_pages = extract_pdf_pages(cv_path)
    except Exception as exc:
        return {"passed": False, "errors": [f"CV PDF cannot be read: {exc}"], "warnings": []}
    try:
        letter_pages = extract_pdf_pages(cover_letter_path)
    except Exception as exc:
        return {"passed": False, "errors": [f"Cover letter PDF cannot be read: {exc}"], "warnings": []}

    cv_text = "\n".join(cv_pages)
    letter_text = "\n".join(letter_pages)
    combined = f"{cv_text}\n{letter_text}"
    names = f"{cv_path.name} {cover_letter_path.name}"

    for term in INTERNAL_TERMS:
        if _contains_term(combined, term) or _contains_term(names, term):
            errors.append(f"internal or technical term exposed: {term}")
    for claim in UNSUPPORTED_CLAIMS:
        if _contains_term(combined, claim):
            errors.append(f"unsupported claim: {claim}")
    for phrase in GENERIC_AI_PHRASES:
        if phrase in combined.lower():
            errors.append(f"generic machine-like phrase: {phrase}")
    for marker in REQUIRED_CV_MARKERS:
        if marker not in cv_text.lower():
            errors.append(f"missing source experience or section: {marker}")

    if len(cv_pages) not in (1, 2):
        errors.append(f"CV must be one or two pages, found {len(cv_pages)}")
    cv_page_words = [len(page.split()) for page in cv_pages]
    if len(cv_pages) == 2 and min(cv_page_words) < 170:
        errors.append(f"unbalanced two-page CV: page word counts {cv_page_words}")
    if len(letter_pages) != 1:
        errors.append(f"cover letter must be one page, found {len(letter_pages)}")
    if len(cv_text.split()) < 480:
        errors.append("CV content is too thin")
    letter_words = len(letter_text.split())
    if not 170 <= letter_words <= 430:
        errors.append(f"cover letter length is not concise and substantial: {letter_words} words")
    if "application for" not in letter_text.lower():
        errors.append("cover letter has no professional subject line")

    if docx_path:
        document = Document(str(docx_path))
        if document.tables:
            errors.append("CV contains tables that may impair parsing")
        allowed_bullets = _source_bullets()
        for paragraph in document.paragraphs:
            if paragraph.style.name == "CV Bullet":
                value = paragraph.text.lstrip("- ").strip()
                if value not in allowed_bullets:
                    errors.append(f"CV bullet is not traceable to source truth: {value[:80]}")

    if analysis:
        title = analysis["job_title"]
        if title and title.lower() not in letter_text.lower():
            errors.append("cover letter does not identify the analysed role")
        if analysis["seniority"]["senior_title"] and re.search(r"\bsenior\b", cv_text, flags=re.IGNORECASE):
            errors.append("CV presents the candidate as senior for a senior vacancy")
        visible_supported = sum(
            1 for keyword in analysis["supported_keywords"]
            if keyword.lower() in cv_text.lower()
        )
        if analysis["supported_keywords"] and visible_supported < min(2, len(analysis["supported_keywords"])):
            errors.append("too few legitimately supported job criteria are visible in the CV")
        if analysis["must_have_gaps"]:
            warnings.append("The vacancy contains unsupported must-have criteria; score cap applied.")

    fingerprint = hashlib.sha256(cv_path.read_bytes() + cover_letter_path.read_bytes()).hexdigest()
    return {
        "passed": not errors,
        "errors": errors,
        "warnings": warnings,
        "cv_pages": len(cv_pages),
        "cv_page_words": cv_page_words,
        "letter_pages": len(letter_pages),
        "letter_words": letter_words,
        "fingerprint": fingerprint,
        "checks": {
            "source_traceability": not any("source truth" in item for item in errors),
            "ats_parseability": not any("tables" in item or "read" in item for item in errors),
            "human_readability": not any("unbalanced" in item or "length" in item for item in errors),
            "anti_hallucination": not any("unsupported" in item for item in errors),
            "anti_automation_artifacts": not any("internal" in item or "machine-like" in item for item in errors),
        },
    }


def validate_application_documents(cv_path, cover_letter_path, analysis=None, docx_path=None):
    report = build_quality_report(cv_path, cover_letter_path, analysis=analysis, docx_path=docx_path)
    return report["passed"], report["errors"]
